import csv
from io import BytesIO, StringIO

import openpyxl
from xhtml2pdf import pisa


def generate_pdf_bytes(html_content):
    """
    Renders an HTML string into PDF binary bytes using xhtml2pdf.
    """
    result = BytesIO()
    # xhtml2pdf requires bytes string or file-like buffer in UTF-8
    pdf = pisa.pisaDocument(
        BytesIO(html_content.encode("utf-8")), result, encoding="utf-8"
    )
    if not pdf.err:
        return result.getvalue()
    return None


def generate_csv_string(headers, rows):
    """
    Generates a CSV formatted string from headers and data rows.
    """
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    return output.getvalue()


def generate_xlsx_bytes(sheet_title, headers, rows):
    """
    Generates an Excel XLSX file binary stream from headers and rows using openpyxl.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title[:30]  # Excel limits tab names to 31 chars

    # Add header row
    ws.append(headers)

    # Add data rows
    for row in rows:
        ws.append([str(item) if item is not None else "" for item in row])

    # Dynamic column width formatting for better readability
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
